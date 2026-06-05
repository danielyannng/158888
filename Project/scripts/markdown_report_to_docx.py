from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_MD = PROJECT_ROOT / "docs" / "FINAL_REPORT.md"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "Geo_located_NLP_Feedback_System_Report.docx"


def clean_inline(text: str) -> str:
    text = text.replace("**", "")
    text = text.replace("__", "")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text_style(cell, *, bold: bool = False, size: int = 9) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.bold = bold


def add_hyperlink_style(doc: Document) -> None:
    styles = doc.styles
    if "Hyperlink" in styles:
        return
    style = styles.add_style("Hyperlink", 2)
    style.font.color.rgb = RGBColor(5, 99, 193)
    style.font.underline = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.name = "Times New Roman"
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in [
        ("Title", 22),
        ("Heading 1", 16),
        ("Heading 2", 14),
        ("Heading 3", 12),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    add_hyperlink_style(doc)


def add_cover_page(doc: Document, markdown: str) -> int:
    lines = markdown.splitlines()
    title = clean_inline(lines[0].lstrip("#").strip())
    metadata: list[str] = []
    cursor = 1
    while cursor < len(lines):
        line = lines[cursor].strip()
        if line == "---":
            cursor += 1
            break
        if line:
            metadata.append(clean_inline(line))
        cursor += 1

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(22)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    srun = subtitle.add_run("Final Project Report")
    srun.bold = True
    srun.font.name = "Times New Roman"
    srun.font.size = Pt(16)

    for item in metadata:
        mp = doc.add_paragraph()
        mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mrun = mp.add_run(item)
        mrun.font.name = "Times New Roman"
        mrun.font.size = Pt(12)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nrun = note.add_run("Massey University 158.888 Information Technology Research Project")
    nrun.italic = True
    nrun.font.name = "Times New Roman"
    nrun.font.size = Pt(12)

    doc.add_page_break()
    return cursor


def is_table_line(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and is_table_line(lines[i]):
        cells = [clean_inline(cell) for cell in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r"[:\-\s]+", cell) for cell in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    table.autofit = True

    for r_idx, row in enumerate(rows):
        for c_idx in range(col_count):
            cell = table.cell(r_idx, c_idx)
            value = row[c_idx] if c_idx < len(row) else ""
            cell.text = value
            set_cell_text_style(cell, bold=(r_idx == 0), size=9)
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")

    doc.add_paragraph()


def add_image_if_exists(doc: Document, source_line: str) -> None:
    matches = re.finditer(r"((?:Project/)?data/evaluation/[^\s`]+\.png)", source_line)
    for match in matches:
        image_ref = match.group(1)
        image_path = PROJECT_ROOT.parent / image_ref if image_ref.startswith("Project/") else PROJECT_ROOT / image_ref
        if not image_path.exists():
            continue
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(5.9))


def add_code_block(doc: Document, code: list[str]) -> None:
    if not code:
        return
    text = "\n".join(code).rstrip()
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)


def add_paragraph_with_inline_format(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cleaned = clean_inline(text)
    run = paragraph.add_run(cleaned)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def convert_markdown_to_docx() -> None:
    markdown = INPUT_MD.read_text(encoding="utf-8")
    lines = markdown.splitlines()

    doc = Document()
    configure_document(doc)
    start_idx = add_cover_page(doc, markdown)

    in_code = False
    code_lines: list[str] = []
    i = start_idx

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue

        if is_table_line(line):
            rows, next_i = parse_table(lines, i)
            add_table(doc, rows)
            i = next_i
            continue

        if stripped.startswith("#### "):
            doc.add_heading(clean_inline(stripped[5:]), level=3)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(clean_inline(stripped[4:]), level=2)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(clean_inline(stripped[3:]), level=1)
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(clean_inline(stripped[2:]), level=1)
            i += 1
            continue

        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if number_match:
            add_paragraph_with_inline_format(doc, number_match.group(1), style="List Number")
            i += 1
            continue

        if stripped.startswith("- "):
            add_paragraph_with_inline_format(doc, stripped[2:], style="List Bullet")
            i += 1
            continue

        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            run = paragraph.add_run(clean_inline(stripped.lstrip("> ")))
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            i += 1
            continue

        add_paragraph_with_inline_format(doc, stripped)
        add_image_if_exists(doc, stripped)
        i += 1

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)


if __name__ == "__main__":
    convert_markdown_to_docx()
    print(OUTPUT_DOCX)
